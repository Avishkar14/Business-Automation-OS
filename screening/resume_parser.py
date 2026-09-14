"""
Extract plain text from a resume file (PDF, DOCX, or TXT).
"""

from pathlib import Path

from pypdf import PdfReader 
from docx import Document 


class ResumeParseError(Exception):
    """Raised when a resume file can't be parsed."""


def parse_resume(file_path: str) -> str:
    """
    Extract raw text from a resume file.

    Args:
        file_path: Path to a .pdf, .docx, or .txt resume file.

    Returns:
        Extracted plain text (whitespace-normalized).

    Raises:
        ResumeParseError: if the file type is unsupported or extraction fails.
    """
    path = Path(file_path)
    if not path.exists():
        raise ResumeParseError(f"File not found: {file_path}")

    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            text = _parse_pdf(path)
        elif suffix == ".docx":
            text = _parse_docx(path)
        elif suffix == ".txt":
            text = path.read_text(encoding="utf-8", errors="ignore")
        else:
            raise ResumeParseError(f"Unsupported file type: {suffix}")
    except ResumeParseError:
        raise
    except Exception as e:
        raise ResumeParseError(f"Failed to parse {file_path}: {e}") from e

    text = _normalize_whitespace(text) 
    if not text.strip():
        raise ResumeParseError(f"No extractable text found in {file_path}")
    return text


def _parse_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _parse_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _normalize_whitespace(text: str) -> str:
    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line]
    return "\n".join(lines)


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print("Usage: python -m screening.resume_parser <resume_file>")
        sys.exit(1)

    print(parse_resume(sys.argv[1]))


